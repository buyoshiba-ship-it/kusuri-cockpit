"""
update_mhlw_data.py
厚生労働省の公式ページから最新の「医療用医薬品供給状況」Excelおよび
最新の「薬価基準収載品目リスト」Excelを自動ダウンロード・解析し、
Webアプリ（cockpit.html / index.html）用の data.js を自動更新＆差分履歴（diff_history.json）を記録するスクリプト。
"""

import os
import re
import ssl
import json
import urllib.request
import urllib.parse
import openpyxl
import io
import sys

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MHLW_SUPPLY_URL = "https://www.mhlw.go.jp/stf/seisakunitsuite/bunya/kenkou_iryou/iryou/kouhatu-iyaku/04_00003.html"
MHLW_PRICE_PORTAL = "https://www.mhlw.go.jp/stf/seisakunitsuite/bunya/0000078916.html"

DATA_JS_PATH = os.path.join(BASE_DIR, "data.js")
PREV_JSON_PATH = os.path.join(BASE_DIR, "prev_medicine_data.json")
DIFF_HISTORY_PATH = os.path.join(BASE_DIR, "diff_history.json")
DOWNLOAD_SUPPLY_PATH = os.path.join(BASE_DIR, "mhlw_latest_supply.xlsx")
YAKKA_MASTER_PATH = os.path.join(BASE_DIR, "yakka_master.json")

def get_ssl_context():
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    return ctx

def get_headers():
    return {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}

import unicodedata

def parse_mhlw_publish_date(title_str):
    """
    厚労省ページのExcelリンクタイトル（例: 医療用医薬品供給状況（令和８年10月６日現在））
    から公式の公表日付を抽出して「YYYY年M月D日」形式で返す。
    取得できない場合はNoneを返す（勝手な当日の日付は絶対に使わない）。
    """
    if not title_str:
        return None
    norm = unicodedata.normalize('NFKC', str(title_str))
    
    # 令和の元号表記: 令和X年Y月Z日
    m_reiwa = re.search(r'令和\s*(\d+)\s*年\s*(\d+)\s*月\s*(\d+)\s*日', norm)
    if m_reiwa:
        reiwa_y = int(m_reiwa.group(1))
        seireki_y = 2018 + reiwa_y
        m = int(m_reiwa.group(2))
        d = int(m_reiwa.group(3))
        return f"{seireki_y}年{m}月{d}日"
        
    # 西暦表記: YYYY年M月D日
    m_seireki = re.search(r'(\d{4})\s*年\s*(\d+)\s*月\s*(\d+)\s*日', norm)
    if m_seireki:
        return f"{m_seireki.group(1)}年{int(m_seireki.group(2))}月{int(m_seireki.group(3))}日"
        
    return None

def fetch_latest_supply_excel_url():
    print("🔍 [供給データ] 厚生労働省の最新ページを確認中...")
    req = urllib.request.Request(MHLW_SUPPLY_URL, headers=get_headers())
    ctx = get_ssl_context()
    
    with urllib.request.urlopen(req, context=ctx) as res:
        html = res.read().decode('utf-8', errors='ignore')
    
    matches = re.findall(r'<a[^>]+href=["\']([^"\']+\.xlsx)["\'][^>]*>(.*?)</a>', html, re.DOTALL | re.IGNORECASE)
    if not matches:
        raise Exception("厚労省供給状況ページからExcelファイルのリンクが見つかりませんでした。")
    
    for href, title in matches:
        clean_title = re.sub(r'<[^>]+>', '', title).strip()
        if "医療用医薬品供給状況" in clean_title or "kyoukyu" in href.lower() or "iyakuhin" in href.lower():
            full_url = urllib.parse.urljoin(MHLW_SUPPLY_URL, href)
            publish_date = parse_mhlw_publish_date(clean_title)
            print(f"✅ 最新供給Excelを発見: {clean_title}")
            print(f"📅 厚労省公式公表日: {publish_date}")
            print(f"🔗 URL: {full_url}")
            return full_url, clean_title, publish_date

    # マッチしなかった場合は最初のxlsxを採用
    href, title = matches[0]
    clean_title = title.strip()
    full_url = urllib.parse.urljoin(MHLW_SUPPLY_URL, href)
    publish_date = parse_mhlw_publish_date(clean_title)
    print(f"✅ 供給Excelを発見: {clean_title} ({full_url})")
    print(f"📅 厚労省公式公表日: {publish_date}")
    return full_url, clean_title, publish_date

def download_supply_excel(url):
    print("📥 最新供給状況Excelをダウンロード中...")
    req = urllib.request.Request(url, headers=get_headers())
    ctx = get_ssl_context()
    with urllib.request.urlopen(req, context=ctx) as res, open(DOWNLOAD_SUPPLY_PATH, 'wb') as f:
        f.write(res.read())
    size_mb = os.path.getsize(DOWNLOAD_SUPPLY_PATH) / (1024 * 1024)
    print(f"✅ ダウンロード完了: {size_mb:.2f} MB")

def fetch_latest_yakka_dictionary():
    """
    厚生労働省の薬価基準収載品目ポータルから最新のExcel（内用薬、注射薬、外用薬、歯科用薬）を取得し、
    YJコード -> 薬価(float) の辞書を作成して返す。
    取得に失敗した場合は既存の yakka_master.json をフォールバックとして使用。
    """
    print("💊 [薬価データ] 厚生労働省の薬価基準収載品目ポータルを確認中...")
    ctx = get_ssl_context()
    try:
        req = urllib.request.Request(MHLW_PRICE_PORTAL, headers=get_headers())
        with urllib.request.urlopen(req, context=ctx) as resp:
            html = resp.read().decode('utf-8', errors='ignore')

        topic_links = re.findall(r'<a[^>]+href=["\']([^"\']+/topics/[^"\']+\.html)["\'][^>]*>(.*?)</a>', html, re.I | re.S)
        target_page_url = None
        for href, text in topic_links:
            clean_text = re.sub(r'<[^>]+>', '', text).strip()
            if "適用" in clean_text or "薬価基準収載品目リスト" in clean_text:
                target_page_url = urllib.parse.urljoin(MHLW_PRICE_PORTAL, href)
                print(f"  🎯 最新薬価告示ページを発見: {clean_text} ({target_page_url})")
                break

        if not target_page_url and topic_links:
            target_page_url = urllib.parse.urljoin(MHLW_PRICE_PORTAL, topic_links[0][0])

        if not target_page_url:
            raise Exception("薬価詳細ページのリンクが見つかりませんでした。")

        req_page = urllib.request.Request(target_page_url, headers=get_headers())
        with urllib.request.urlopen(req_page, context=ctx) as resp:
            page_html = resp.read().decode('utf-8', errors='ignore')

        excel_links = re.findall(r'<a[^>]+href=["\']([^"\']+\.xlsx)["\'][^>]*>(.*?)</a>', page_html, re.I | re.S)
        category_files = {}
        for href, text in excel_links:
            full_href = urllib.parse.urljoin(target_page_url, href)
            m = re.search(r'[-_](0[1-4])\.xlsx$', href, re.I)
            if m:
                cat_num = m.group(1)
                if cat_num not in category_files:
                    category_files[cat_num] = full_href

        if not category_files:
            raise Exception("薬価Excelファイルリンクが抽出できませんでした。")

        yakka_dict = {}
        for cat, url in sorted(category_files.items()):
            req_file = urllib.request.Request(url, headers=get_headers())
            with urllib.request.urlopen(req_file, context=ctx) as resp:
                wb = openpyxl.load_workbook(io.BytesIO(resp.read()), read_only=True)
                sheet = wb.active
                
                yj_col = 1
                price_col = 12
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i == 0:
                        for col_idx, val in enumerate(row):
                            s = str(val or "")
                            if "コード" in s:
                                yj_col = col_idx
                            elif "薬価" in s and "基準" not in s:
                                price_col = col_idx
                        continue
                    if not row or len(row) <= max(yj_col, price_col):
                        continue
                    yj = str(row[yj_col] or "").strip()
                    raw_price = row[price_col]
                    if not yj or raw_price is None:
                        continue
                    try:
                        yakka_dict[yj] = round(float(raw_price), 2)
                    except (ValueError, TypeError):
                        pass

        print(f"✅ 厚労省薬価リスト解析完了: 合計 {len(yakka_dict):,} 件の薬価を同期")
        with open(YAKKA_MASTER_PATH, 'w', encoding='utf-8') as f:
            json.dump(yakka_dict, f, ensure_ascii=False)
        return yakka_dict

    except Exception as e:
        print(f"⚠️ 薬価の最新ダウンロード・解析中にエラー（既存マスターを使用します）: {e}")
        if os.path.exists(YAKKA_MASTER_PATH):
            with open(YAKKA_MASTER_PATH, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {}

def parse_excel_to_dataset(excel_path, yakka_dict):
    print("⚙️ 供給状況Excelデータを解析・構造化中...")
    wb = openpyxl.load_workbook(excel_path, read_only=True)
    ws = wb.active

    # 統一名収載（一般名収載）対策：同規格プレフィックス（YJ先頭9桁）の薬価フォールバックマップ構築
    prefix9_map = {}
    prefix9_groups = {}
    for yj, p in yakka_dict.items():
        if len(yj) >= 9:
            p9 = yj[:9]
            if p9 not in prefix9_groups:
                prefix9_groups[p9] = []
            prefix9_groups[p9].append((yj, p))

    for p9, items in prefix9_groups.items():
        # 統一名収載コード（末尾01X番台）が存在する場合はそれを優先
        touitsu = [p for yj, p in items if len(yj) >= 12 and yj[9:11] == '01']
        if touitsu:
            prefix9_map[p9] = touitsu[0]
        else:
            # 存在しない場合はグループ内の代表薬価を採用
            prefix9_map[p9] = items[0][1]

    medicine_data = []
    
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i <= 1 or not row or len(row) < 14:
            continue
        
        name = str(row[5] or "").strip()
        ingredient = str(row[2] or "").strip()
        maker = str(row[6] or "").strip()
        yj_code = str(row[4] or "").strip()
        status_str = str(row[11] or "").strip()
        reason = str(row[13] or "").strip()
        recovery = str(row[14] or "").strip()
        category = str(row[9] or "").strip()
        
        # 更新日: col 12 (出荷対応更新日) または col 19 (その他更新日)
        update_date = ""
        if row[12] is not None and str(row[12]).strip() not in ["None", ""]:
            update_date = str(row[12]).strip()
        elif len(row) > 19 and row[19] is not None and str(row[19]).strip() not in ["None", ""]:
            update_date = str(row[19]).strip()
            
        ship_vol = str(row[16] or "").strip() if len(row) > 16 else ""
        
        # 製品区分・基礎的医薬品タグ
        raw_prod = str(row[7] or "")
        raw_basic = str(row[8] or "")
        
        tags = []
        if "先発品" in raw_prod:
            tags.append("先")
        elif "準先発品" in raw_prod:
            tags.append("準先")
        elif "後発品" in raw_prod:
            tags.append("後")
        elif "長期収載品" in raw_prod:
            tags.append("旧後")
            
        if "1：対象" in raw_basic or "対象" in raw_basic:
            tags.append("基")
            
        prod_tag = " ".join(tags)
        
        # 解除見込み時期 (col 15: 2027年1月など)
        rec_timing = str(row[15] or "").strip() if len(row) > 15 and row[15] is not None else ""
        if rec_timing in ["None", "－", "-"]:
            rec_timing = ""
            
        # 今回更新NEWマーク (col 20: 'New' の場合)
        is_new = 1 if len(row) > 20 and str(row[20] or "").strip() == "New" else 0
            
        # ステータス区分: 0=通常, 1=限定出荷, 2=供給停止
        st_type = 0
        if "供給停止" in status_str or "停止" in status_str:
            st_type = 2
        elif "限定出荷" in status_str:
            st_type = 1
        else:
            st_type = 0
            
        # 薬価 (個別YJコード ➔ なければ統一名収載の同規格プレフィックスで補完)
        price = yakka_dict.get(yj_code, None)
        if price is None and len(yj_code) >= 9:
            price = prefix9_map.get(yj_code[:9], None)

        medicine_data.append([
            name,         # 0
            ingredient,   # 1
            maker,        # 2
            yj_code,      # 3
            status_str,   # 4
            reason,       # 5
            recovery,     # 6
            st_type,      # 7
            category,     # 8
            update_date,  # 9
            ship_vol,     # 10
            prod_tag,     # 11
            rec_timing,   # 12
            is_new,       # 13
            price         # 14 (薬価 float または None)
        ])
        
    print(f"✅ 全 {len(medicine_data):,} 品目の抽出に成功しました！")
    matched_price_count = sum(1 for m in medicine_data if m[14] is not None)
    print(f"💰 薬価紐付け完了: {matched_price_count:,} 品目 ({matched_price_count/len(medicine_data)*100:.1f}%)")
    return medicine_data

def detect_diff_and_save(new_data):
    """前回データと比較して差分（改善・悪化・停止）を検知"""
    prev_map = {}
    if os.path.exists(PREV_JSON_PATH):
        try:
            with open(PREV_JSON_PATH, "r", encoding="utf-8") as f:
                prev_list = json.load(f)
                for item in prev_list:
                    yj = item[3]
                    if yj:
                        prev_map[yj] = item
        except Exception as e:
            print("⚠️ 前回データの読み込みスキップ:", e)

    diffs = {
        "improved": [],
        "worsened": [],
        "new_stopped": [],
        "total_new_items": 0
    }
    
    for row in new_data:
        yj = row[3]
        if not yj or yj not in prev_map:
            diffs["total_new_items"] += 1
            continue
            
        prev_row = prev_map[yj]
        prev_st = prev_row[7] # 0, 1, 2
        new_st = row[7]
        
        name = row[0]
        maker = row[2]
        
        # 改善（例: 限定出荷/停止 ➔ 通常、停止 ➔ 限定）
        if new_st < prev_st:
            diffs["improved"].append({
                "yj": yj,
                "name": name,
                "maker": maker,
                "from_status": prev_row[4],
                "to_status": row[4],
                "changeType": "improved"
            })
        # 悪化（通常 ➔ 限定/停止、限定 ➔ 停止）
        elif new_st > prev_st:
            is_stopped = (new_st == 2)
            diff_entry = {
                "yj": yj,
                "name": name,
                "maker": maker,
                "from_status": prev_row[4],
                "to_status": row[4],
                "changeType": "worsened"
            }
            diffs["worsened"].append(diff_entry)
            if is_stopped:
                diffs["new_stopped"].append(diff_entry)
                
    print(f"\n📊 差分検知結果:")
    print(f"  🟢 改善・通常復帰した品目: {len(diffs['improved'])} 件")
    print(f"  🔴 悪化・出荷制限になった品目: {len(diffs['worsened'])} 件 (うち供給停止: {len(diffs['new_stopped'])} 件)")
    
    # 差分履歴を保存
    with open(DIFF_HISTORY_PATH, "w", encoding="utf-8") as f:
        json.dump(diffs, f, ensure_ascii=False, indent=2)
        
    # 今回のデータを次回比較用として保存
    with open(PREV_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(new_data, f, ensure_ascii=False)

def export_data_js(medicine_data, publish_date=None):
    print("💾 data.js を出力中...")
    json_str = json.dumps(medicine_data, ensure_ascii=False)
    
    date_str_val = f'"{publish_date}"' if publish_date else 'null'
    js_content = f"const MHLW_PUBLISH_DATE = {date_str_val};\nconst MEDICINE_DATA = {json_str};\n"
    
    with open(DATA_JS_PATH, "w", encoding="utf-8") as f:
        f.write(js_content)
        
    size_mb = os.path.getsize(DATA_JS_PATH) / (1024 * 1024)
    print(f"🎉 data.js の更新が完了しました！（公表日: {publish_date} / ファイルサイズ: {size_mb:.2f} MB）")

def main():
    print("=" * 60)
    print(" 🏥 厚生労働省 医薬品供給＆薬価データ 自動同期パイプライン")
    print("=" * 60)
    
    try:
        # 1. 薬価マスターの取得・更新
        yakka_dict = fetch_latest_yakka_dictionary()

        # 2. 供給状況Excelの取得（厚労省ページの公式公表日を正確にトレース）
        excel_url, title, publish_date = fetch_latest_supply_excel_url()
        download_supply_excel(excel_url)

        # 3. 供給状況と薬価の結合解析
        dataset = parse_excel_to_dataset(DOWNLOAD_SUPPLY_PATH, yakka_dict)

        # 4. 差分検知＆履歴保存
        detect_diff_and_save(dataset)

        # 5. data.js 出力（厚労省の公式公表日を埋め込み）
        export_data_js(dataset, publish_date)
        print("\n✨ すべての工程が正常に完了しました！")
    except Exception as e:
        print(f"\n❌ エラーが発生しました: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()
